"""
logsentry.reporters.excel_reporter — 4-sheet Excel workbook.
"""
from __future__ import annotations

from typing import TYPE_CHECKING

from ..analyzers.engine import AnalysisReport, Severity

if TYPE_CHECKING:
    pass


# Color palette (HEX without #)
COLOR = {
    Severity.CRITICAL: "C0392B",
    Severity.HIGH:     "E67E22",
    Severity.MEDIUM:   "F1C40F",
    Severity.LOW:      "5DADE2",
    Severity.INFO:     "AAB7B8",
    Severity.CLEAN:    "27AE60",
}
HEADER_BG = "1F2937"
HEADER_FG = "FFFFFF"


class ExcelReporter:
    def __init__(self, report: AnalysisReport) -> None:
        self.report = report

    def generate(self, path: str) -> str:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        wb.remove(wb.active)

        thin = Side(style="thin", color="D5D8DC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_font = Font(bold=True, color=HEADER_FG, size=11)
        header_fill = PatternFill("solid", fgColor=HEADER_BG)
        wrap = Alignment(wrap_text=True, vertical="top")
        center = Alignment(horizontal="center", vertical="center")

        def style_header(ws, row: int, ncols: int) -> None:
            for c in range(1, ncols + 1):
                cell = ws.cell(row=row, column=c)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = border

        # ---------------- Summary ----------------
        ws = wb.create_sheet("📊 Summary")
        s = self.report.summary()
        ws["A1"] = "🛰️  LOGSENTRY — Linux Log Forensics"
        ws["A1"].font = Font(bold=True, size=16, color="1F2937")
        ws.merge_cells("A1:D1")

        rows = [
            ("Host",      s["hostname"]),
            ("OS",        s["os"]),
            ("Kernel",    s["kernel"]),
            ("Collected", s["timestamp"]),
            ("",          ""),
            ("Threat Score", f"{s['score']}/100"),
            ("Grade",     s["grade"]),
            ("Detectors", s["total_detectors"]),
            ("Clean",     s["clean"]),
            ("Triggered", s["triggered"]),
        ]
        for i, (k, v) in enumerate(rows, start=3):
            ws.cell(row=i, column=1, value=k).font = Font(bold=True)
            ws.cell(row=i, column=2, value=v)

        # severity table
        ws.cell(row=14, column=1, value="Severity").font = Font(bold=True)
        ws.cell(row=14, column=2, value="Count").font = Font(bold=True)
        style_header(ws, 14, 2)
        sev_rows = [
            ("CRITICAL", s["critical"], COLOR[Severity.CRITICAL]),
            ("HIGH",     s["high"],     COLOR[Severity.HIGH]),
            ("MEDIUM",   s["medium"],   COLOR[Severity.MEDIUM]),
            ("LOW",      s["low"],      COLOR[Severity.LOW]),
        ]
        for i, (name, count, color) in enumerate(sev_rows, start=15):
            cell = ws.cell(row=i, column=1, value=name)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=True, color="FFFFFF")
            ws.cell(row=i, column=2, value=count)

        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 40

        # ---------------- Events ----------------
        ws = wb.create_sheet("🚨 Events")
        headers = ["ID", "Category", "Title", "Severity", "Status",
                   "Detail", "Recommendation", "Reference", "IOCs"]
        ws.append(headers)
        style_header(ws, 1, len(headers))
        for r in self.report.results:
            ws.append([
                r.check_id, r.category, r.title, r.severity.value, r.status,
                r.detail, r.recommendation, r.reference, ", ".join(r.iocs),
            ])
            row = ws.max_row
            color = COLOR.get(r.severity, "FFFFFF")
            sev_cell = ws.cell(row=row, column=4)
            sev_cell.fill = PatternFill("solid", fgColor=color)
            sev_cell.font = Font(bold=True, color="FFFFFF")
            sev_cell.alignment = center
            for c in range(1, len(headers) + 1):
                ws.cell(row=row, column=c).alignment = wrap
                ws.cell(row=row, column=c).border = border

        widths = [10, 12, 28, 10, 8, 50, 40, 16, 30]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"

        # ---------------- By Category ----------------
        ws = wb.create_sheet("📂 By Category")
        ws.append(["Category", "Total", "Clean", "Triggered", "Trigger rate"])
        style_header(ws, 1, 5)
        for cat, items in self.report.by_category().items():
            total = len(items)
            triggered = sum(1 for r in items if r.is_issue)
            clean = total - triggered
            row = ws.max_row + 1
            ws.cell(row=row, column=1, value=cat).font = Font(bold=True)
            ws.cell(row=row, column=2, value=total)
            ws.cell(row=row, column=3, value=clean)
            ws.cell(row=row, column=4, value=triggered)
            ws.cell(row=row, column=5,
                    value=f"=IF(B{row}=0,0,D{row}/B{row})").number_format = "0.0%"
        for col, w in zip("ABCDE", (16, 8, 8, 12, 14)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

        # ---------------- IOCs ----------------
        ws = wb.create_sheet("🎯 IOCs")
        ws.append(["Indicator", "First seen in", "Severity"])
        style_header(ws, 1, 3)
        seen = {}
        for r in self.report.results:
            for ioc in r.iocs:
                if ioc not in seen:
                    seen[ioc] = (r.title, r.severity.value)
        for ioc, (where, sev) in seen.items():
            ws.append([ioc, where, sev])
        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 32
        ws.column_dimensions["C"].width = 12
        ws.freeze_panes = "A2"

        wb.save(path)
        return path
